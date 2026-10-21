import io
import re
from collections.abc import Iterable

from openpyxl import Workbook, load_workbook
from sqlalchemy import Select, func, or_, select
from sqlalchemy.orm import Session, selectinload

from app import models, schemas


PROVINCE_BY_CITY = {
    "上海": "上海",
    "南京": "江苏",
    "苏州": "江苏",
    "无锡": "江苏",
    "常州": "江苏",
    "杭州": "浙江",
    "宁波": "浙江",
    "合肥": "安徽",
    "芜湖": "安徽",
    "武汉": "湖北",
}


def company_query(
    q: str | None = None,
    province: str | None = None,
    industry: str | None = None,
    icp_tier: str | None = None,
    stage: models.OpportunityStage | None = None,
    rd_size_min: int | None = None,
) -> Select:
    statement = select(models.Company).options(selectinload(models.Company.aliases))
    if q:
        statement = statement.outerjoin(models.CompanyAlias).where(
            or_(models.Company.name.ilike(f"%{q}%"), models.CompanyAlias.alias.ilike(f"%{q}%"))
        )
    if province:
        statement = statement.where(models.Company.province == province)
    if industry:
        statement = statement.where(models.Company.industry == industry)
    if icp_tier:
        statement = statement.where(models.Company.icp_tier == icp_tier)
    if stage:
        statement = statement.where(models.Company.current_stage == stage)
    if rd_size_min is not None:
        statement = statement.where(models.Company.rd_size >= rd_size_min)
    return statement.distinct()


def parse_rd_size(value: object) -> int | None:
    if value is None:
        return None
    text = str(value).strip()
    numbers = [int(x) for x in re.findall(r"\d+", text)]
    return numbers[0] if numbers else None


def infer_region(location: str) -> tuple[str, str | None]:
    for city, province in PROVINCE_BY_CITY.items():
        if city in location:
            return province, city
    return "待核", location or None


def normalize_universe_row(headers: list[str], values: Iterable[object]) -> dict | None:
    row = dict(zip(headers, values))
    name = str(row.get("企业名称") or row.get("name") or "").strip()
    if not name:
        return None
    location = str(row.get("总部/华东研发地") or row.get("city") or "").strip()
    province, city = infer_region(location)
    priority = str(row.get("销售优先级") or row.get("icp_tier") or "待评估").strip()
    icp_tier = "A" if "★★★★★" in priority else "B" if "★★★★" in priority else "待评估"
    return {
        "name": name,
        "industry": str(row.get("一级行业") or row.get("industry") or "待分类").strip(),
        "province": str(row.get("province") or province).strip(),
        "city": city,
        "rd_size": parse_rd_size(row.get("软件/嵌入式研发规模(估)") or row.get("rd_size")),
        "icp_tier": icp_tier,
        "crm_status": str(row.get("crm_status") or "未建档").strip(),
        "sales_owner": row.get("sales_owner"),
    }


def import_companies(session: Session, content: bytes) -> schemas.ImportResult:
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    created = skipped = 0
    errors: list[str] = []
    ignored_sheets = {"口径说明", "汇总", "README", "说明"}
    for sheet in workbook.worksheets:
        if sheet.title in ignored_sheets:
            continue
        rows = sheet.iter_rows(values_only=True)
        headers = [str(value).strip() if value is not None else "" for value in next(rows, ())]
        if "企业名称" not in headers and "name" not in headers:
            continue
        for row_number, values in enumerate(rows, start=2):
            try:
                normalized = normalize_universe_row(headers, values)
                if not normalized:
                    continue
                exists = session.scalar(select(models.Company.id).where(models.Company.name == normalized["name"]))
                if exists:
                    skipped += 1
                    continue
                session.add(models.Company(**normalized))
                created += 1
            except Exception as exc:  # preserve row-level import progress
                errors.append(f"{sheet.title}!{row_number}: {exc}")
    session.commit()
    return schemas.ImportResult(created=created, skipped=skipped, errors=errors)


def export_companies(companies: list[models.Company]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "企业"
    sheet.append(["企业名称", "一级行业", "省份", "城市", "研发规模", "ICP分层", "CRM状态", "销售Owner", "机会状态", "模拟数据"])
    for company in companies:
        sheet.append([
            company.name,
            company.industry,
            company.province,
            company.city,
            company.rd_size,
            company.icp_tier,
            company.crm_status,
            company.sales_owner,
            company.current_stage.value,
            "是" if company.is_sample else "否",
        ])
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def dashboard_summary(session: Session) -> schemas.DashboardSummary:
    stage_rows = session.execute(
        select(models.Company.current_stage, func.count()).group_by(models.Company.current_stage)
    ).all()
    province_rows = session.execute(
        select(models.Company.province, func.count()).group_by(models.Company.province)
    ).all()
    return schemas.DashboardSummary(
        company_count=session.scalar(select(func.count(models.Company.id))) or 0,
        event_count=session.scalar(select(func.count(models.Event.id))) or 0,
        verified_event_count=session.scalar(
            select(func.count(models.Event.id)).where(models.Event.verification_status == "已核验")
        ) or 0,
        stage_counts={stage.value: count for stage, count in stage_rows},
        province_counts={province: count for province, count in province_rows},
        sample_data=bool(session.scalar(select(func.count(models.Company.id)).where(models.Company.is_sample.is_(True)))),
    )

